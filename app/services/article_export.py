from io import BytesIO
from datetime import date
from decimal import Decimal
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_articles(items):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Artikel'
    sheet.append(['Marke', 'Bezeichnung', 'Art. Nr.', 'Lief. Art. Nr.', 'EAN', 'Farbe', 'Grösse', 'Geliefert gesamt', 'Einheit', 'Letzter UVP (CHF)', 'UVP vom', 'Erste Lieferung', 'Letzte Lieferung'])
    for item in items:
        # Separate rows preserve incompatible units instead of adding pieces and pairs.
        for delivered in item['delivered'] or [{'quantity': None, 'unit': None}]:
            values = [item[k] for k in ['brand', 'description', 'article_no', 'supplier_article_no', 'ean', 'color', 'size']]
            values += [Decimal(delivered['quantity']) if delivered['quantity'] is not None else None, delivered['unit'], Decimal(item['latest_uvp']) if item['latest_uvp'] is not None else None, item['uvp_date'], item['first_seen'], item['last_seen']]
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Keep identifiers and formula-like text literal.
                    cell.number_format = '@'
            sheet.cell(sheet.max_row, 8).number_format = '0.00'
            sheet.cell(sheet.max_row, 10).number_format = '0.00 "CHF"'
            for column in [11, 12, 13]:
                sheet.cell(sheet.max_row, column).number_format = 'dd.mm.yyyy'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='303238')
    for index, width in enumerate([20, 48, 20, 22, 22, 30, 16, 20, 12, 22, 18, 18, 18], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal='left')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="Artikel_{date.today().isoformat()}.xlsx"'})
