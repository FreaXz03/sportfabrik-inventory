from io import BytesIO
from openpyxl import load_workbook
from test_catalog import client


def test_export_filters_all_pages_and_cell_types(client):
    response = client.get('/api/articles/export?page_size=1&page=2&brand=Hoka')
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 3  # Hoka has deliveries in two different units.
    assert sheet['D2'].value == '0012345678901'
    assert sheet['D2'].data_type == 's'
    assert sheet['G2'].data_type == 'n'
    assert sheet['I2'].value == 200
    assert sheet['J2'].is_date
    assert {sheet['H2'].value, sheet['H3'].value} == {'PAA', 'STK'}
    assert sheet.freeze_panes == 'A2'
    assert load_workbook(BytesIO(client.get('/api/articles/export?brand=Missing').content)).active.max_row == 1


def test_export_formula_text_stays_text(client):
    from app.core.database import get_session
    from app.core.models import Artikel
    generator = client.app.dependency_overrides[get_session]()
    session = next(generator)
    try:
        session.get(Artikel, 1).bezeichnung = '=1+1'
        session.commit()
    finally:
        generator.close()
    sheet = load_workbook(BytesIO(client.get('/api/articles/export?brand=Hoka').content)).active
    assert sheet['B2'].value == '=1+1'
    assert sheet['B2'].data_type == 's'
