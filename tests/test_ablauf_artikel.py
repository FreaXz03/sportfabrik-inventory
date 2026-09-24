"""Ablauf „Artikel nachschlagen": suchen, sortieren, blättern, exportieren,
Lieferhistorie und Preisverlauf ansehen, Notizen schreiben; Etiketten für
einen ganzen Wareneingang drucken.

Deckt die Artikelsuche (Inbox vom 23.09.2026), die Artikeldetails und F3
(alle dürfen alles lesen) über die echte App mit Anmeldung ab.
"""

from io import BytesIO

from conftest import ANNA, BEAT, CHEF
from openpyxl import load_workbook
from sqlalchemy import select
from testbelege import POSITIONEN, importieren, kopf, rechnung_pdf

from app.core.models import Wareneingang


def test_suchen_exportieren_historie_preise_notizen(welt):
    client = welt.client
    welt.anmelden(CHEF)
    assert importieren(client, rechnung_pdf(header_lines=kopf(nummer="9000000001", datum="05.02.2026"))).status_code == 200
    # Nachlieferung Poloshirt M mit neuem UVP, dazu eine Formel als Bezeichnung.
    nachlieferung = [POSITIONEN[0][:8] + ["54.90", "33.00"], ["Hoka", "324100", "H1", "1", "0012345678905", "=1+1", "1", "Paa", "150.00", "90.00"]]
    assert importieren(client, rechnung_pdf(header_lines=kopf(nummer="9000000002", datum="05.08.2026"), rows=nachlieferung)).status_code == 200

    welt.anmelden(ANNA)
    alle = client.get("/api/articles").json()
    assert alle["total"] == 4
    assert client.get("/api/brands").json()
    polo = client.get("/api/articles?ean=4006632041234").json()["items"][0]
    assert polo["brand"] == "Nike"
    # Suche, Filter, Sortierung, Seiten.
    assert client.get("/api/articles?q=Laufschuh").json()["total"] == 1
    assert client.get("/api/articles?supplier_article_no=A1").json()["total"] == 2
    assert client.get("/api/articles?last_delivery_from=2026-08-01").json()["total"] == 2
    assert client.get("/api/articles?last_delivery_from=kein-datum").status_code == 422
    marken = [i["brand"] for i in client.get("/api/articles?sort_by=brand&sort_dir=desc").json()["items"]]
    assert marken == sorted(marken, reverse=True)
    assert client.get("/api/articles?sort_by=passwort").status_code == 422
    seite2 = client.get("/api/articles?page_size=3&page=2").json()
    assert len(seite2["items"]) == 1 and seite2["total"] == 4

    # Export als Excel: alle Seiten, EAN als Text, Formeln bleiben Text.
    export = client.get("/api/articles/export?page_size=1&brand=Hoka")
    assert export.status_code == 200
    blatt = load_workbook(BytesIO(export.content)).active
    assert blatt.max_row == 2 and blatt.freeze_panes == "A2"
    assert blatt["D2"].value == "0012345678905" and blatt["D2"].data_type == "s"
    assert blatt["B2"].value == "=1+1" and blatt["B2"].data_type == "s"

    # Lieferhistorie und Preisverlauf (chronologisch).
    historie = client.get(f"/api/articles/{polo['id']}/history").json()
    assert historie["product"]["manuell"] is False
    # Alle Lieferungen desselben Artikels (beide Grössen), diese Variante zweimal.
    assert [i["ean"] for i in historie["items"]].count("4006632041234") == 2
    preise = client.get(f"/api/articles/{polo['id']}/prices").json()["items"]
    assert [p["uvp"] for p in preise] == ["49.90", "54.90"]
    assert client.get("/api/articles/999999/history").status_code == 404

    # Notizen: Anna schreibt, Beat darf lesen, aber nicht ändern; der
    # Filialleiter darf; veraltete Version wird abgelehnt.
    url = f"/api/articles/{polo['id']}/notes"
    for leer in ("", "   ", "x" * 2001):
        assert client.post(url, json={"body": leer}).status_code == 422
    notiz = client.post(url, json={"body": "  Für Stammkundin reservieren  "}).json()
    assert notiz["body"] == "Für Stammkundin reservieren" and notiz["author_name"] == "Anna"
    ziel = f"{url}/{notiz['id']}"
    assert client.put(ziel, json={"body": "Neu", "version": 1}).json()["version"] == 2
    assert client.put(ziel, json={"body": "Veraltet", "version": 1}).status_code == 409
    welt.anmelden(BEAT)
    assert client.get(url).json()["items"][0]["can_edit"] is False
    assert client.put(ziel, json={"body": "Fremd", "version": 2}).status_code == 403
    assert client.delete(f"{ziel}?version=2").status_code == 403
    welt.anmelden(CHEF)
    bearbeitet = client.put(ziel, json={"body": "Freigegeben", "version": 2}).json()
    assert bearbeitet["updated_by"] == "Chef" and bearbeitet["author_name"] == "Anna"
    assert client.delete(f"{ziel}?version=2").status_code == 409
    assert client.delete(f"{ziel}?version=3").status_code == 204
    assert client.get(url).json()["total"] == 0

    # Etiketten für einen ganzen Wareneingang, ein PDF.
    with welt.sessions() as session:
        wareneingang_id = session.scalar(select(Wareneingang.id).order_by(Wareneingang.id))
    etiketten = client.get(f"/api/wareneingaenge/{wareneingang_id}/etiketten.pdf")
    assert etiketten.status_code == 200 and etiketten.content.startswith(b"%PDF")

    # Seiten der Artikelansicht.
    for pfad in ("/articles", f"/articles/{polo['id']}/history", "/invoices"):
        assert client.get(pfad).status_code == 200, pfad
